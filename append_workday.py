import json
import openpyxl
import os

wd_file = "/Users/iamrj846/Desktop/iamrj846.github.io/resources/workday_companies.json"
excel_file = "/Users/iamrj846/Desktop/iamrj846.github.io/resources/job_urls.xlsx"

# Load workday companies
with open(wd_file, "r") as f:
    wd_data = json.load(f)

print(f"Loaded {len(wd_data)} Workday companies from JSON")

# Load excel
wb = openpyxl.load_workbook(excel_file)
sheet_name = wb.sheetnames[0]
for name in wb.sheetnames:
    if "Master ATS Directory" in name:
        sheet_name = name
        break
ws = wb[sheet_name]

# Check existing rows to prevent duplicates
existing_urls = set()
for row in ws.iter_rows(min_row=2, values_only=True):
    if row and len(row) > 2 and row[2]:
        existing_urls.add(row[2])

print(f"Current rows in Excel: {ws.max_row}")

added = 0
for item in wd_data:
    parts = item.split('|')
    if len(parts) == 3:
        comp, wd_domain, path = parts
        endpoint = f"https://{comp}.{wd_domain}.myworkdayjobs.com/wday/cxs/{comp}/{path}/jobs"
        apply_base = f"https://{comp}.{wd_domain}.myworkdayjobs.com/en-US/{path}"
        
        # some company names might just be the slug capitalized
        comp_name = comp.capitalize()
        
        if endpoint not in existing_urls:
            ws.append((comp_name, "Workday", endpoint, apply_base))
            added += 1

print(f"Adding {added} Workday endpoints to Excel...")
wb.save(excel_file)
print("Done saving Excel.")

# Remove unused json files
unused = [
    "/Users/iamrj846/Desktop/iamrj846.github.io/resources/icims_companies.json",
    "/Users/iamrj846/Desktop/iamrj846.github.io/resources/paylocity_companies_clean.json"
]
for p in unused:
    if os.path.exists(p):
        os.remove(p)
        print(f"Removed unused file: {p}")

