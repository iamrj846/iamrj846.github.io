import openpyxl

file_path = "/Users/iamrj846/Desktop/iamrj846.github.io/resources/job_urls.xlsx"
wb = openpyxl.load_workbook(file_path)

sheet_name = wb.sheetnames[0]
for name in wb.sheetnames:
    if "Master ATS Directory" in name:
        sheet_name = name
        break

ws = wb[sheet_name]

# Oracle companies to add
oracle_companies = [
    ("JPMorgan Chase", "Oracle", "https://jpmc.fa.oraclecloud.com/hcmRestApi/resources/latest/recruitingCEJobRequisitions?onlyData=true&expand=requisitionList&finder=findReqs;siteNumber=CX_1001", "https://jpmc.fa.oraclecloud.com/hcmUI/CandidateExperience/en/sites/CX_1001/"),
    ("Marriott", "Oracle", "https://marriott.fa.oraclecloud.com/hcmRestApi/resources/latest/recruitingCEJobRequisitions?onlyData=true&expand=requisitionList&finder=findReqs;siteNumber=CX_1", "https://marriott.fa.oraclecloud.com/hcmUI/CandidateExperience/en/sites/CX_1/")
]

for row in oracle_companies:
    ws.append(row)

wb.save(file_path)
print("Successfully appended Oracle ATS companies to job_urls.xlsx")
